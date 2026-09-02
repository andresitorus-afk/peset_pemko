import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json

df = pd.read_csv("ml/hasil_v2.csv")

def safe_int(val, default=0):
    try:
        if pd.isna(val):
            return default
        return int(val)
    except (ValueError, TypeError):
        return default

wb = Workbook()
ws = wb.active
ws.title = "Hasil Prediksi Aset"

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Calibri", size=10)
TITLE_FONT = Font(name="Calibri", bold=True, size=14, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="666666")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

REDAH_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
SEDANG_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
TINGGI_FILL = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
ALT_ROW_FILL = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")

MAPPING_KIB = {0: "Tanah", 1: "Peralatan", 2: "Gedung"}
MAPPING_KONDISI = {0: "Baik", 1: "Rusak Ringan", 2: "Rusak Berat"}
MAPPING_STATUS = {0: "Tidak Produktif", 1: "Produktif", 2: "Sangat Produktif"}

poi_cols = [c for c in df.columns if c.startswith("poi_")]

ws.merge_cells("A1:Q1")
ws["A1"] = "DATA PREDIKSI PENGELOLAAN ASET PEMERINTAH KOTA MEDAN"
ws["A1"].font = TITLE_FONT
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

ws.merge_cells("A2:Q2")
ws["A2"] = f"Total Data: {len(df)} aset | Sumber: hasil_v2.csv"
ws["A2"].font = SUBTITLE_FONT
ws["A2"].alignment = Alignment(horizontal="center")
ws.row_dimensions[2].height = 20

header_row = 4
headers = [
    "No", "ID Aset", "OPD", "Kategori KIB", "Kondisi", "Status",
    "Luas (m2)", "Umur (thn)", "Alamat",
    "Jenis Aktual", "Kontribusi (Rp)", "Permintaan Aktual",
    "Prediksi Jenis", "Prob. Prediksi", "Prediksi Kontribusi", "Prediksi Permintaan", "Prob. Permintaan"
]

for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=header_row, column=col_idx, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER

ws.row_dimensions[header_row].height = 35

for row_idx, (_, row) in enumerate(df.iterrows(), 1):
    excel_row = header_row + row_idx
    is_alt = row_idx % 2 == 0

    prob_jenis_val = ""
    try:
        pj = json.loads(str(row["prob_jenis"]).replace("'", '"'))
        prob_jenis_val = ", ".join(f"{k}: {v:.1%}" for k, v in pj.items())
    except:
        prob_jenis_val = str(row.get("prob_jenis", ""))

    prob_perm_val = ""
    try:
        pp = json.loads(str(row["prob_permintaan"]).replace("'", '"'))
        prob_perm_val = ", ".join(f"{k}: {v:.1%}" for k, v in pp.items())
    except:
        prob_perm_val = str(row.get("prob_permintaan", ""))

    kib_label = MAPPING_KIB.get(safe_int(row["kategori_kib"]), str(row["kategori_kib"]))
    kondisi_label = MAPPING_KONDISI.get(safe_int(row["kondisi"]), str(row["kondisi"]))
    status_label = MAPPING_STATUS.get(safe_int(row["status"]), str(row["status"]))

    pred_perm_raw = row.get("prediksi_permintaan", "")
    pred_perm = str(pred_perm_raw) if pd.isna(pred_perm_raw) else str(pred_perm_raw)

    values = [
        row_idx,
        str(row.get("aset_id", "")),
        safe_int(row["opd_id"]),
        kib_label,
        kondisi_label,
        status_label,
        safe_int(row["luas_m2"]),
        safe_int(row["umur_aset"]),
        str(row.get("address", "")),
        str(row.get("y_jenis", "")),
        safe_int(row["y_kontribusi"]),
        safe_int(row["y_permintaan"]),
        str(row.get("prediksi_jenis", "")),
        prob_jenis_val,
        safe_int(row["prediksi_kontribusi_rp"]),
        pred_perm,
        prob_perm_val,
    ]

    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=excel_row, column=col_idx, value=val)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=(col_idx == 9))

        if is_alt:
            cell.fill = ALT_ROW_FILL

        if col_idx in (11, 15):
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="right", vertical="center")

        if col_idx == 7:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal="right", vertical="center")

        if col_idx == 16:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    pred_jenis = str(row.get("prediksi_jenis", ""))
    prob_jenis_map = {}
    try:
        prob_jenis_map = json.loads(str(row["prob_jenis"]).replace("'", '"'))
    except:
        pass
    max_prob = max(prob_jenis_map.values()) if prob_jenis_map else 0

    if pred_jenis in ["KSP", "KSPI"]:
        status_cell = ws.cell(row=excel_row, column=13)
        if max_prob > 0.8:
            status_cell.fill = TINGGI_FILL
        elif max_prob > 0.5:
            status_cell.fill = SEDANG_FILL
        else:
            status_cell.fill = REDAH_FILL

    pred_perm_label = str(row.get("prediksi_permintaan", "")).lower()
    if pred_perm_label == "tinggi":
        ws.cell(row=excel_row, column=16).fill = TINGGI_FILL
    elif pred_perm_label == "sedang":
        ws.cell(row=excel_row, column=16).fill = SEDANG_FILL
    elif pred_perm_label == "rendah":
        ws.cell(row=excel_row, column=16).fill = REDAH_FILL

    cond_val = safe_int(row["kondisi"])
    cond_cell = ws.cell(row=excel_row, column=5)
    if cond_val == 2:
        cond_cell.fill = TINGGI_FILL
        cond_cell.font = Font(name="Calibri", size=10, color="C62828", bold=True)
    elif cond_val == 1:
        cond_cell.fill = SEDANG_FILL

    status_val = safe_int(row["status"])
    stat_cell = ws.cell(row=excel_row, column=6)
    if status_val == 0:
        stat_cell.fill = REDAH_FILL

col_widths = {
    1: 5, 2: 12, 3: 6, 4: 14, 5: 16, 6: 18,
    7: 10, 8: 10, 9: 55,
    10: 12, 11: 18, 12: 14,
    13: 14, 14: 45, 15: 20, 16: 16, 17: 45,
}
for col, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col)].width = width

ws.auto_filter.ref = f"A{header_row}:Q{header_row + len(df)}"
ws.freeze_panes = f"A{header_row + 1}"
ws.sheet_properties.tabColor = "1F4E79"

ws2 = wb.create_sheet("Statistik Ringkas")

stat_data = [
    ["METRIK", "NILAI"],
    ["Total Aset", len(df)],
    ["", ""],
    ["DISTRIBUSI KATEGORI KIB", ""],
    ["Tanah", len(df[df["kategori_kib"] == 0])],
    ["Peralatan", len(df[df["kategori_kib"] == 1])],
    ["Gedung", len(df[df["kategori_kib"] == 2])],
    ["", ""],
    ["DISTRIBUSI KONDISI", ""],
    ["Baik", len(df[df["kondisi"] == 0])],
    ["Rusak Ringan", len(df[df["kondisi"] == 1])],
    ["Rusak Berat", len(df[df["kondisi"] == 2])],
    ["", ""],
    ["DISTRIBUSI STATUS", ""],
    ["Tidak Produktif", len(df[df["status"] == 0])],
    ["Produktif", len(df[df["status"] == 1])],
    ["Sangat Produktif", len(df[df["status"] == 2])],
    ["", ""],
    ["RATA-RATA", ""],
    ["Luas (m2)", f"{df['luas_m2'].mean():,.0f}"],
    ["Umur (tahun)", f"{df['umur_aset'].mean():,.1f}"],
    ["Kontribusi (Rp)", f"Rp {df['y_kontribusi'].mean():,.0f}"],
    ["", ""],
    ["PREDIKSI JENIS", ""],
    ["SEWA", len(df[df["prediksi_jenis"] == "SEWA"])],
    ["PKP", len(df[df["prediksi_jenis"] == "PKP"])],
    ["KSP", len(df[df["prediksi_jenis"] == "KSP"])],
    ["BGS", len(df[df["prediksi_jenis"] == "BGS"])],
    ["BSG", len(df[df["prediksi_jenis"] == "BSG"])],
    ["KSPI", len(df[df["prediksi_jenis"] == "KSPI"])],
    ["", ""],
    ["PREDIKSI PERMINTAAN", ""],
    ["Rendah", len(df[df["prediksi_permintaan"] == "rendah"])],
    ["Sedang", len(df[df["prediksi_permintaan"] == "sedang"])],
    ["Tinggi", len(df[df["prediksi_permintaan"] == "tinggi"])],
]

for r_idx, row_data in enumerate(stat_data, 1):
    for c_idx, val in enumerate(row_data, 1):
        cell = ws2.cell(row=r_idx, column=c_idx, value=val)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER

        if r_idx == 1:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
        elif c_idx == 1 and r_idx > 1 and row_data[1] == "":
            cell.font = Font(name="Calibri", bold=True, size=10, color="1F4E79")

ws2.column_dimensions["A"].width = 25
ws2.column_dimensions["B"].width = 20
ws2.sheet_properties.tabColor = "2E75B6"

ws3 = wb.create_sheet("Rincian POI")

poi_headers = ["ID Aset", "Alamat"] + poi_cols
for c_idx, h in enumerate(poi_headers, 1):
    cell = ws3.cell(row=1, column=c_idx, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center", wrap_text=True)

ws3.row_dimensions[1].height = 30

for r_idx, (_, row) in enumerate(df.iterrows(), 2):
    ws3.cell(row=r_idx, column=1, value=str(row.get("aset_id", ""))).font = DATA_FONT
    ws3.cell(row=r_idx, column=2, value=str(row.get("address", ""))).font = DATA_FONT
    ws3.cell(row=r_idx, column=2).alignment = Alignment(wrap_text=True)

    for c_idx, col in enumerate(poi_cols, 3):
        val = row[col]
        cell = ws3.cell(row=r_idx, column=c_idx, value=val)
        cell.font = DATA_FONT
        cell.border = THIN_BORDER

        if "_n" in col:
            cell.number_format = '0'
            cell.alignment = Alignment(horizontal="center")
        elif "_d" in col:
            cell.number_format = '0.0000'
            cell.alignment = Alignment(horizontal="center")

    if r_idx % 2 == 0:
        for c_idx in range(1, len(poi_headers) + 1):
            ws3.cell(row=r_idx, column=c_idx).fill = ALT_ROW_FILL

poi_col_widths = {1: 12, 2: 55}
for i in range(3, len(poi_headers) + 1):
    poi_col_widths[i] = 14
for col, width in poi_col_widths.items():
    ws3.column_dimensions[get_column_letter(col)].width = width

ws3.auto_filter.ref = f"A1:{get_column_letter(len(poi_headers))}1"
ws3.freeze_panes = "C2"
ws3.sheet_properties.tabColor = "548235"

output_path = "ml/hasil_v2_formatted.xlsx"
wb.save(output_path)
print(f"Excel saved: {output_path}")
print(f"Sheets: {wb.sheetnames}")
print(f"Rows: {len(df)}")
